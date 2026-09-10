"""Export regionale "A01" (Regione Veneto, art. 49 c.4 L.R. 21/2008): elenco infortuni per Provincia,
una riga per persona. Anteprima di qualità del dato prima della generazione (docs/00 Q3-Q4).

Tracciato (dal file di riferimento): n., Data, Orario (08-11 | 11-14 | 14-17 | 17-08), Sesso (M | F),
Età (0-10 … oltre 80), Pista (Denominazione, Codice regionale), Attrezzatura (Sci | Snowboard | Altro),
Tipologia (Caduta accidentale | Collisione altro sciatore | Collisione ostacolo fisso | Collisione ostacolo
mobile | Malore | Incidente impianto risalita | Altro), Soccorso congiunto.
"""

from __future__ import annotations

import datetime as dt
import io
from dataclasses import dataclass, field
from typing import Any

from django.db.models import QuerySet

from safe.apps.rescue.models import Event, Person
from safe.apps.rescue.views import PERSON_PREFETCH, PERSON_SELECT, scope_events
from safe.apps.territory.models import IstatAdminUnit

from .common import age_class, local
from .duplicates import DuplicateGroup, find_duplicates

TEMPLATE = "veneto_a01"
TIME_SLOTS = ("08,00 - 11,00", "11,00 - 14,00", "14,00 - 17,00", "17,00 - 08,00")
AGE_CLASSES = ("0-10", "11-20", "21-30", "31-40", "41-50", "51-60", "61-70", "71-80", "oltre 80")
EQUIPMENT = ("Sci", "Snowboard", "Altro")
TYPOLOGY = (
    "Caduta accidentale",
    "Collisione altro sciatore",
    "Collisione ostacolo fisso",
    "Collisione ostacolo mobile",
    "Malore",
    "Incidente impianto risalita",
    "Altro",
)
JOINT_COL = "Soccorsi effettuati unitamente ad altro personale"


@dataclass
class Selection:
    season: str
    area: IstatAdminUnit | None
    persons: list[Person]
    excluded_buildings: list[str] = field(default_factory=list)
    excluded_outside: list[str] = field(default_factory=list)
    no_geometry: list[str] = field(default_factory=list)
    groups: list[DuplicateGroup] = field(default_factory=list)
    slopes_without_code: set[str] = field(default_factory=set)
    merge_duplicates: bool = True


def select(
    user,
    season: str,
    area: IstatAdminUnit | None,
    teams: list[str] | None,
    *,  # noqa: ANN001
    exclude_buildings: bool = True,
    merge_duplicates: bool = True,
) -> Selection:
    events: QuerySet = scope_events(Event.objects.filter(deleted_at__isnull=True, season__code=season), user)
    if teams:
        events = events.filter(team_id__in=teams)
    events = events.select_related("team", "slope", "location_type", "cause")
    sel = Selection(season=season, area=area, persons=[], merge_duplicates=merge_duplicates)
    kept_ids = []
    for e in events:
        if exclude_buildings and e.location_type and e.location_type.mapping.get("regional_export_exclude"):
            sel.excluded_buildings.append(str(e.id))
            continue
        if area is not None:
            if e.geom is None:
                sel.no_geometry.append(str(e.id))
            elif not area.geom.contains(e.geom):
                sel.excluded_outside.append(str(e.id))
                continue
        if e.slope and not e.slope.regional_code:
            sel.slopes_without_code.add(str(e.slope_id))
        kept_ids.append(e.id)
    persons = list(
        Person.objects.filter(event_id__in=kept_ids, deleted_at__isnull=True)
        .select_related(*PERSON_SELECT)
        .prefetch_related(*PERSON_PREFETCH)
        .order_by("event__dateandtime", "sequence")
    )
    sel.persons = persons
    if merge_duplicates and persons:
        company = persons[0].company
        sel.groups = find_duplicates(persons, company.duplicate_rule, company.timezone)
    return sel


def preview(sel: Selection) -> dict[str, Any]:
    merged = sum(len(g.merged_ids) for g in sel.groups)
    return {
        "persons_total": len(sel.persons) + 0,
        "excluded_in_buildings": {
            "count": len(sel.excluded_buildings),
            "event_ids": sel.excluded_buildings[:200],
        },
        "excluded_outside_area": {
            "count": len(sel.excluded_outside),
            "event_ids": sel.excluded_outside[:200],
        },
        "duplicate_groups": {
            "count": len(sel.groups),
            "persons_merged": merged,
            "groups": [g.as_dict() for g in sel.groups[:200]],
        },
        "exportable_persons": len(sel.persons) - (merged if sel.merge_duplicates else 0),
        "warnings": (
            (
                [
                    {
                        "code": "slope_missing_regional_code",
                        "count": len(sel.slopes_without_code),
                        "slope_ids": sorted(sel.slopes_without_code),
                    }
                ]
                if sel.slopes_without_code
                else []
            )
            + (
                [
                    {
                        "code": "events_without_geometry",
                        "count": len(sel.no_geometry),
                        "event_ids": sel.no_geometry[:200],
                    }
                ]
                if sel.no_geometry
                else []
            )
            + (
                [{"code": "administrative_boundaries_not_loaded"}]
                if sel.area is None and not IstatAdminUnit.objects.exists()
                else []
            )
        ),
    }


def rows(sel: Selection, tz: str) -> list[list[Any]]:
    merged_ids = {pid for g in sel.groups for pid in g.merged_ids} if sel.merge_duplicates else set()
    joint_ids = {g.kept_id for g in sel.groups if g.joint} if sel.merge_duplicates else set()
    out: list[list[Any]] = []
    n = 0
    for p in sel.persons:
        if str(p.id) in merged_ids:
            continue
        n += 1
        e = p.event
        d = local(e.dateandtime, tz)
        hour = d.hour + d.minute / 60
        slot = 0 if 8 <= hour < 11 else 1 if 11 <= hour < 14 else 2 if 14 <= hour < 17 else 3
        gender = p.gender.code if p.gender else None
        cls = age_class(p.age, "veneto_a01")
        cls_label = "oltre 80" if cls == "80+" else cls
        equip = p.equipment.mapping.get(TEMPLATE, "Altro") if p.equipment else ""
        typ = e.cause.mapping.get(TEMPLATE, "Altro") if e.cause else ""
        row: list[Any] = [n, d.strftime("%d/%m/%Y")]
        row += [1 if slot == i else "" for i in range(4)]
        row += [1 if gender == "male" else "", 1 if gender == "female" else ""]
        row += [1 if cls_label == c else "" for c in AGE_CLASSES]
        row += [e.slope.name if e.slope else "NON RILEVATO", (e.slope.regional_code if e.slope else "") or ""]
        row += [1 if equip == c else "" for c in EQUIPMENT]
        row += [1 if typ == c else "" for c in TYPOLOGY]
        row += [1 if str(p.id) in joint_ids else ""]
        out.append(row)
    return out


HEADER_TOP = [
    "n.",
    "Data",
    "Orario",
    "",
    "",
    "",
    "Sesso",
    "",
    "Età",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "",
    "Pista",
    "",
    "Attrezzatura",
    "",
    "",
    "Tipologia",
    "",
    "",
    "",
    "",
    "",
    "",
    JOINT_COL,
]
HEADER_SUB = [
    "",
    "",
    *TIME_SLOTS,
    "Maschio",
    "Femmina",
    *AGE_CLASSES,
    "Denominazione",
    "Codice regionale",
    *EQUIPMENT,
    *TYPOLOGY,
    "",
]


def build_xlsx(sel: Selection, company_name: str, tz: str) -> bytes:
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font

    wb = Workbook()
    ws = wb.active
    ws.title = "Foglio1"
    ws.append(["Elenco infortuni verificatisi sulle piste da sci del Veneto"])
    ws.append([f"Stagione invernale {sel.season}"])
    ws.append([])
    ws.append([f"Società/Ente: {company_name}"])
    ws.append([])
    ws.append(HEADER_TOP)
    ws.append([])
    ws.append(HEADER_SUB)
    for r in rows(sel, tz):
        ws.append(r)
    ws["A1"].font = Font(bold=True, size=12)
    for cell in ws[6]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center")
    for cell in ws[8]:
        cell.alignment = Alignment(horizontal="center", wrap_text=True)
    ws.freeze_panes = "C9"
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def build_xls(sel: Selection, company_name: str, tz: str) -> bytes:
    import xlwt

    wb = xlwt.Workbook(encoding="utf-8")
    ws = wb.add_sheet("Foglio1")
    bold = xlwt.easyxf("font: bold on")
    ws.write(0, 0, "Elenco infortuni verificatisi sulle piste da sci del Veneto", bold)
    ws.write(1, 0, f"Stagione invernale {sel.season}")
    ws.write(3, 0, f"Società/Ente: {company_name}")
    for c, v in enumerate(HEADER_TOP):
        if v:
            ws.write(5, c, v, bold)
    for c, v in enumerate(HEADER_SUB):
        if v:
            ws.write(7, c, v)
    for r, row in enumerate(rows(sel, tz), start=8):
        for c, v in enumerate(row):
            if v != "":
                ws.write(r, c, v)
    out = io.BytesIO()
    wb.save(out)
    return out.getvalue()


def filename(sel: Selection, fmt: str) -> str:
    area = f"_{sel.area.code}" if sel.area else ""
    return f"A01_{sel.season.replace('/', '-')}{area}.{fmt}"


_ = dt
