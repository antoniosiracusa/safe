"""Export dataset, PDF del rapporto, export regionale A01 con anteprima e duplicati."""

import base64
import datetime as dt
import io

import pytest
from django.contrib.gis.geos import MultiPolygon, Polygon
from django.template.loader import render_to_string
from openpyxl import load_workbook

from safe.apps.jobs.models import AsyncJob
from safe.apps.org.models import Team
from safe.apps.rescue import services
from safe.apps.rescue.models import Event, PersonEvacuationMean
from safe.apps.tenancy.context import tenant_context
from safe.apps.territory.models import IstatAdminUnit, SkiArea

from . import factories as f
from .conftest import auth, create_user

pytestmark = pytest.mark.django_db
UTC = dt.UTC


@pytest.fixture
def dataset(company, team):  # noqa: ANN001, ANN201
    """E1 (Squadra A, 18/01 10:05 UTC = 11:05 locali) con P1 F 26 sci e P2 M 16 snowboard;
    E2 (Carabinieri, 18/01 10:20 UTC) con P3 F 26 sci → duplicato di P1;
    E3 in edificio (escluso); E4 16:30 UTC = 17:30 locali (fascia 17-08)."""
    with tenant_context(company.id):
        _, zone, slope = f.territory()
        carab = Team.objects.create(name="Carabinieri")
        e1 = f.event(team, when=dt.datetime(2026, 1, 18, 10, 5, tzinfo=UTC))
        f.person(e1)
        f.person(
            e1,
            age=16,
            gender=f.lv("gender", "male"),
            country_id="PL",
            equipment=f.lv("equipment", "snowboard"),
        )
        e2 = f.event(carab, when=dt.datetime(2026, 1, 18, 10, 20, tzinfo=UTC))
        f.person(e2)
        e3 = f.event(
            team,
            when=dt.datetime(2026, 1, 19, 12, 0, tzinfo=UTC),
            location_type=f.lv("location_type", "building"),
        )
        f.person(e3, age=40)
        e4 = f.event(team, when=dt.datetime(2026, 2, 2, 16, 30, tzinfo=UTC), cause=f.lv("cause", "illness"))
        p4 = f.person(e4, age=82, gender=f.lv("gender", "male"))
        PersonEvacuationMean.objects.create(
            person=p4, mean=f.lv("evacuation_mean", "helicopter_118"), order=2
        )
        for ev in Event.objects.all():
            services.compute_event_validity(ev)
        return {"e1": e1, "e2": e2, "e3": e3, "e4": e4, "carab": carab}


# --- dataset ---------------------------------------------------------------------------------------


def test_dataset_export_xlsx_events(as_admin, dataset):
    r = as_admin.post(
        "/api/v1/exports/dataset",
        {"dataset": "events", "format": "xlsx", "filters": {"season": "2025/2026"}},
        format="json",
    )
    assert r.status_code == 202
    job = r.json()
    assert job["status"] == "done" and job["download_url"].endswith("/download")  # eager in test
    d = as_admin.get(job["download_url"])
    assert d.status_code == 200 and d["Content-Type"].startswith("application/vnd.openxmlformats")
    ws = load_workbook(io.BytesIO(d.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0] == "ID evento" and len(rows) == 5
    assert rows[1][0] == str(dataset["e4"].id)[-8:].upper()  # ordinati per data decrescente
    assert rows[1][3] == "17:30" and "Malore" in rows[1] and rows[1][22] == 1


def test_dataset_export_csv_persons_has_no_names(as_admin, dataset):
    r = as_admin.post(
        "/api/v1/exports/dataset", {"dataset": "persons", "format": "csv", "lang": "en"}, format="json"
    )
    d = as_admin.get(r.json()["download_url"])
    text = d.content.decode("utf-8-sig")
    assert text.startswith("Event ID;Season;Date;Time")
    assert text.count("\r\n") == 6  # header + 5 persone
    assert "firstname" not in text.lower() and "surname initial" in text.lower()


def test_exports_list_and_permissions(as_admin, api_client, company, team, dataset):
    as_admin.post("/api/v1/exports/dataset", {"dataset": "events", "format": "csv"}, format="json")
    body = as_admin.get("/api/v1/exports").json()
    assert body["count"] == 1 and body["results"][0]["export"]["template"] == "dataset_events"
    user = create_user(company, "noexp@a.test", roles=("rescuer",), team=team)
    assert (
        auth(api_client, user)
        .post("/api/v1/exports/dataset", {"dataset": "events"}, format="json")
        .status_code
        == 403
    )


# --- PDF -------------------------------------------------------------------------------------------


def test_event_report_pdf(as_admin, dataset, company):
    e1 = dataset["e1"]
    r = as_admin.get(f"/api/v1/events/{e1.id}/report.pdf")
    assert r.status_code == 200 and r["Content-Type"] == "application/pdf"
    assert r.content[:5] == b"%PDF-"
    assert r["Content-Disposition"].startswith('attachment; filename="report-evento-2026-01-18-')
    # seconda richiesta: stesso job in cache
    with tenant_context(company.id):
        assert AsyncJob.objects.filter(kind="pdf_report").count() == 1
    as_admin.get(f"/api/v1/events/{e1.id}/report.pdf")
    with tenant_context(company.id):
        assert AsyncJob.objects.filter(kind="pdf_report").count() == 1
    # modifica → nuovo PDF
    as_admin.patch(f"/api/v1/events/{e1.id}", {"note": "aggiornato"}, format="json")
    as_admin.get(f"/api/v1/events/{e1.id}/report.pdf")
    with tenant_context(company.id):
        assert AsyncJob.objects.filter(kind="pdf_report").count() == 2


def test_report_pdf_brand_logo_and_name(dataset, company, settings, tmp_path):
    """Il PDF usa nome e logo della società di servizio (BRAND_NAME / BRAND_LOGO_URL) come la SPA."""
    from safe.apps.reports import pdf

    png = base64.b64decode(
        "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNkYPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
    )
    (tmp_path / "ski-civetta.png").write_bytes(png)
    settings.BRAND_NAME = "SAFECIVETTA.IT"
    settings.BRAND_LOGO_URL = "brand/ski-civetta.png"
    settings.BRAND_DIR = str(tmp_path)
    with tenant_context(company.id):
        ctx = pdf.build_context(dataset["e1"], "it", downloaded_by="test", with_map=False)
    assert ctx["brand_name"] == "SAFECIVETTA.IT"
    assert ctx["brand_logo"].startswith("data:image/png;base64,")
    html = render_to_string("reports/event_report.html", ctx)
    assert 'class="logo"' in html and "SAFECIVETTA.IT" in html and 'class="mark"' not in html
    # senza logo configurato resta il simbolo predefinito
    settings.BRAND_LOGO_URL = ""
    assert pdf.brand_context()["brand_logo"] is None


def test_person_conditions_in_pdf_and_export(as_admin, dataset, company):
    """Condizioni generali e primo soccorso prestato compaiono nelle griglie del PDF e nell'export persone."""
    from safe.apps.reports import pdf
    from safe.apps.rescue.models import PersonCondition, PersonFirstAid

    e1 = dataset["e1"]
    with tenant_context(company.id):
        p = e1.persons.order_by("sequence").first()
        PersonCondition.objects.create(person=p, condition=f.lv("condition", "conscious"))
        PersonFirstAid.objects.create(person=p, first_aid=f.lv("first_aid", "transport"))
        ctx = pdf.build_context(e1, "it", downloaded_by="test", with_map=False)
    pc = ctx["persons"][0]
    assert [c["label"] for c in pc["condition_grid"] if c["checked"]] == ["Cosciente"]
    assert [c["label"] for c in pc["first_aid_grid"] if c["checked"]] == ["Trasporto"]
    r = as_admin.post("/api/v1/exports/dataset", {"dataset": "persons", "format": "csv"}, format="json")
    text = as_admin.get(r.json()["download_url"]).content.decode("utf-8-sig")
    header, *rows = [ln for ln in text.splitlines() if ln]
    cols = header.split(";")
    i, j = cols.index("Condizioni generali"), cols.index("Primo soccorso prestato")
    assert any(row.split(";")[i] == "Cosciente" and row.split(";")[j] == "Trasporto" for row in rows)


def test_report_pdf_requires_permission(api_client, company, team, dataset):
    user = create_user(company, "nopdf@a.test", roles=("key_custodian",), team=team)
    r = auth(api_client, user).get(f"/api/v1/events/{dataset['e1'].id}/report.pdf")
    assert r.status_code == 403 and r.json()["missing"] == ["reports.pdf"]


# --- regionale A01 -------------------------------------------------------------------------------


def test_regional_preview_counts(as_admin, dataset):
    r = as_admin.post(
        "/api/v1/exports/regional/preview", {"template": "veneto_a01", "season": "2025/2026"}, format="json"
    )
    assert r.status_code == 200, r.content
    pv = r.json()
    assert pv["persons_total"] == 4  # P1, P2, P3, P4 (E3 in edificio escluso)
    assert pv["excluded_in_buildings"]["count"] == 1
    assert pv["duplicate_groups"]["count"] == 1 and pv["duplicate_groups"]["persons_merged"] == 1
    grp = pv["duplicate_groups"]["groups"][0]
    assert sorted(grp["teams"]) == ["Carabinieri", "Squadra A"] and grp["joint"] is True
    assert pv["exportable_persons"] == 3
    assert any(w["code"] == "administrative_boundaries_not_loaded" for w in pv["warnings"])


def test_regional_export_a01_rows(as_admin, dataset):
    r = as_admin.post("/api/v1/exports/regional", {"season": "2025/2026", "format": "xlsx"}, format="json")
    assert r.status_code == 202
    job = r.json()
    assert job["status"] == "done" and job["result"]["preview"]["exportable_persons"] == 3
    d = as_admin.get(job["download_url"])
    assert d["Content-Disposition"] == 'attachment; filename="A01_2025-2026.xlsx"'
    ws = load_workbook(io.BytesIO(d.content)).active
    rows = list(ws.iter_rows(values_only=True))
    assert rows[0][0].startswith("Elenco infortuni") and rows[1][0] == "Stagione invernale 2025/2026"
    assert rows[7][2] == "08,00 - 11,00" and rows[7][17] == "Denominazione"
    data = rows[8:]
    assert len(data) == 3
    first = data[0]  # P1: 11:05 locali → fascia 11-14, F, 21-30, LE CIAUNE C.1.24, sci, collisione, congiunto
    assert first[0] == 1 and first[1] == "18/01/2026"
    assert first[3] == 1 and first[7] == 1 and first[10] == 1
    assert (
        first[17] == "LE CIAUNE"
        and first[18] == "C.1.24"
        and first[19] == 1
        and first[23] == 1
        and first[29] == 1
    )
    last = data[-1]  # P4: 17:30 locali → fascia 17-08, M, oltre 80, malore
    assert last[5] == 1 and last[6] == 1 and last[16] == 1 and last[26] == 1 and last[29] is None


def test_regional_export_xls_and_area_filter(as_admin, dataset, company):
    IstatAdminUnit.objects.create(
        level="province",
        code="025",
        name="Belluno",
        edition_year=2025,
        geom=MultiPolygon(Polygon(((12.0, 46.4), (12.2, 46.4), (12.2, 46.5), (12.0, 46.5), (12.0, 46.4)))),
    )
    IstatAdminUnit.objects.create(
        level="province",
        code="999",
        name="Altrove",
        edition_year=2025,
        geom=MultiPolygon(Polygon(((10.0, 44.0), (10.1, 44.0), (10.1, 44.1), (10.0, 44.1), (10.0, 44.0)))),
    )
    areas = as_admin.get("/api/v1/exports/administrative-areas").json()
    assert areas["loaded"] is True and sorted(a["code"] for a in areas["areas"]) == ["025", "999"]
    # con il confine del comprensorio compaiono solo le aree intersecate, comuni inclusi
    IstatAdminUnit.objects.create(
        level="municipality",
        code="025059",
        name="Val di Zoldo",
        edition_year=2025,
        geom=MultiPolygon(
            Polygon(((12.05, 46.42), (12.15, 46.42), (12.15, 46.48), (12.05, 46.48), (12.05, 46.42)))
        ),
    )
    with tenant_context(company.id):
        area = SkiArea.objects.get(name="Civetta")
        area.boundary = MultiPolygon(
            Polygon(((12.08, 46.43), (12.12, 46.43), (12.12, 46.46), (12.08, 46.46), (12.08, 46.43)))
        )
        area.save(update_fields=["boundary"])
    areas = as_admin.get("/api/v1/exports/administrative-areas").json()["areas"]
    assert [(a["level"], a["name"]) for a in areas] == [
        ("province", "Belluno"),
        ("municipality", "Val di Zoldo"),
    ]
    pv = as_admin.post(
        "/api/v1/exports/regional/preview",
        {"season": "2025/2026", "administrative_area": {"level": "province", "code": "999"}},
        format="json",
    ).json()
    assert pv["excluded_outside_area"]["count"] == 3 and pv["exportable_persons"] == 0
    r = as_admin.post(
        "/api/v1/exports/regional",
        {"season": "2025/2026", "format": "xls", "administrative_area": {"level": "province", "code": "025"}},
        format="json",
    )
    job = r.json()
    assert job["status"] == "done" and job["result_filename"] == "A01_2025-2026_025.xls"
    d = as_admin.get(job["download_url"])
    assert d["Content-Type"] == "application/vnd.ms-excel" and d.content[:4] == b"\xd0\xcf\x11\xe0"
